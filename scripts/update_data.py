"""
Ambil data terbaru dari Google Drive dan simpan sebagai JSON di folder data/.
Dijalankan otomatis oleh GitHub Actions (.github/workflows/update-data.yml).

Butuh environment variable GOOGLE_SERVICE_ACCOUNT_JSON berisi isi file JSON
service account Google yang sudah diberi akses Viewer ke kedua spreadsheet
dan folder GSC (lihat README.md).
"""
import csv
import datetime as dt
import io
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CONFIG = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GSHEET = "application/vnd.google-apps.spreadsheet"

MONTHS = {"jan": 1, "januari": 1, "january": 1, "feb": 2, "februari": 2, "february": 2, "mar": 3, "maret": 3, "march": 3,
          "apr": 4, "april": 4, "mei": 5, "may": 5, "jun": 6, "juni": 6, "june": 6, "jul": 7, "juli": 7, "july": 7,
          "agu": 8, "agus": 8, "agustus": 8, "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
          "okt": 10, "oktober": 10, "oct": 10, "october": 10, "nov": 11, "november": 11,
          "des": 12, "desember": 12, "dec": 12, "december": 12}
METRICS = {"# review": "reviews", "campaign redeem": "redeem", "rating": "rating", "traffic": "traffic",
           "directions": "directions", "call": "calls", "chat whatsapp": "chats", "webiste click": "web", "website click": "web"}
GSC_DOMAINS = [("ruheediary.com", "Ruheediary"), ("ruhee.id", "Ruhee"), ("laverelash.com", "Lavere"), ("ever.co.id", "Ever")]
GSC_LIMIT = 200


def month_key(s):
    m = re.match(r"^\s*([A-Za-z]+)\s+(\d{4})\s*$", str(s or ""))
    if not m or m.group(1).lower() not in MONTHS:
        return None
    return f"{m.group(2)}-{MONTHS[m.group(1).lower()]:02d}"


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v == v else None
    if isinstance(v, str):
        s = v.strip().replace(",", ".")
        if re.match(r"^-?\d+(\.\d+)?$", s):
            return float(s)
    return None


def clean(x):
    return int(x) if isinstance(x, float) and x.is_integer() else x


# ---------- parsers (same rules as the Claude dashboard loaders) ----------
def parse_monitoring(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    names = {ws.title.strip(): ws for ws in wb.worksheets}
    rows = []
    for sheet in ["EVER", "RUHEE", "LAVERE"]:
        ws = names.get("Monitoring GMB " + sheet)
        if ws is None:
            raise RuntimeError(f'Sheet "Monitoring GMB {sheet}" tidak ditemukan')
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(grid) if r and r[0] == "Column 1"), None)
        if hi is None:
            raise RuntimeError(f"Header 'Column 1' tidak ditemukan di Monitoring GMB {sheet}")
        branches = []
        for c in range(2, len(grid[hi])):
            h = str(grid[hi][c] or "").strip()
            if re.match(r"^kpi", h, re.I):
                break
            if h:
                branches.append((c, h))
        metric = None
        for row in grid[hi + 1:]:
            if not row:
                continue
            label = str(row[0] or "").strip().lower()
            if label in METRICS:
                metric = METRICS[label]
            month = month_key(row[1] if len(row) > 1 else None)
            if not metric or not month:
                continue
            for c, branch in branches:
                v = num(row[c]) if c < len(row) else None
                if v is not None:
                    rows.append({"sheet": sheet, "branch": branch, "metric": metric, "month": month, "value": clean(v)})
    return rows


def parse_reviews(csv_text):
    grid = list(csv.reader(io.StringIO(csv_text)))
    top, sub = grid[0], grid[1]
    blocks, c = [], 0
    while c < len(sub):
        if re.match(r"^WEEK 1 \(P\)$", sub[c].strip(), re.I):
            end = c
            while end < len(sub) and not re.match(r"^positif$", sub[end].strip(), re.I):
                end += 1
            month = None
            for k in range(c, min(end + 1, len(top))):
                month = month or month_key(top[k])
            weeks = []
            for k in range(c, end):
                m = re.match(r"^WEEK (\d) \(([PN])\)$", sub[k].strip(), re.I)
                if m:
                    weeks.append((k, int(m.group(1)), m.group(2).upper()))
            if month:
                blocks.append((month, weeks))
            c = end
        c += 1
    out = []
    for r in grid[2:]:
        outlet = (r[0] if r else "").strip()
        if not outlet:
            continue
        for month, weeks in blocks:
            by = {}
            for k, w, pn in weeks:
                s = r[k].strip() if k < len(r) else ""
                if not re.match(r"^-?\d+(\.\d+)?$", s):
                    continue
                by.setdefault(w, {"positive": 0, "negative": 0})["positive" if pn == "P" else "negative"] = clean(float(s))
            for w, v in by.items():
                out.append({"outlet": outlet, "month": month, "week": w, **v})
    return out


def as_day(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)):
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v or "").strip()
    return s[:10] if re.match(r"^\d{4}-\d{2}-\d{2}", s) else None


def gnum(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    s = str(v).strip()
    pct = "%" in s
    try:
        x = float(s.replace("%", "").replace(",", "."))
    except ValueError:
        return None
    return x / 100 if pct else x


def parse_gsc(xlsx_bytes, title, brand, domain):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    sh = {ws.title.strip().lower(): [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    daily = []
    for r in (sh.get("chart") or sh.get("dates") or [])[1:]:
        d = as_day(r[0]) if r else None
        if d:
            daily.append((d, gnum(r[1]), gnum(r[2]), gnum(r[3]), gnum(r[4])))
    if not daily:
        return [], False
    days = sorted(d[0] for d in daily)
    base = {"brand": brand, "domain": domain, "file": title, "start": days[0], "end": days[-1]}
    rows, cut = [], False
    for d, c, i, ctr, p in daily:
        rows.append({**base, "type": "day", "key": d, "clicks": clean(c), "impressions": clean(i), "ctr": ctr, "position": p})
    for sheet, typ, limit in [("queries", "query", GSC_LIMIT), ("pages", "page", GSC_LIMIT), ("devices", "device", 10), ("countries", "country", 15)]:
        g = [r for r in (sh.get(sheet) or [])[1:] if r and r[0] is not None]
        cut = cut or len(g) > limit
        for r in g[:limit]:
            rows.append({**base, "type": typ, "key": str(r[0]), "clicks": clean(gnum(r[1])), "impressions": clean(gnum(r[2])), "ctr": gnum(r[3]), "position": gnum(r[4])})
    return rows, cut


# ---------- Google Drive ----------
def drive():
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    info = json.loads(os.environ["GOOGLE_SERVICE_ACCOUNT_JSON"])
    creds = service_account.Credentials.from_service_account_info(info, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def export(svc, file_id, mime):
    return svc.files().export(fileId=file_id, mimeType=mime).execute()


def download(svc, file_id):
    return svc.files().get_media(fileId=file_id, supportsAllDrives=True).execute()


def write(name, payload):
    DATA.mkdir(exist_ok=True)
    (DATA / name).write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def main():
    svc = drive()
    report = {}

    mon = parse_monitoring(export(svc, CONFIG["monitoring_sheet_id"], XLSX))
    write("gmb_monitoring.json", {"rows": mon})
    report["gmb_monitoring"] = len(mon)

    csv_text = export(svc, CONFIG["review_sheet_id"], "text/csv").decode("utf-8")
    rev = parse_reviews(csv_text)
    write("reviews.json", {"rows": rev})
    report["reviews"] = len(rev)

    q = f"'{CONFIG['gsc_folder_id']}' in parents and trashed = false and (mimeType = '{XLSX}' or mimeType = '{GSHEET}')"
    files, token = [], None
    while True:
        res = svc.files().list(q=q, pageSize=100, pageToken=token, fields="nextPageToken, files(id,name,mimeType,modifiedTime)",
                               supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
        files += res.get("files", [])
        token = res.get("nextPageToken")
        if not token:
            break
    gsc, skipped, cut = [], [], False
    for f in sorted(files, key=lambda f: f.get("modifiedTime", "")):
        hit = next(((d, b) for d, b in GSC_DOMAINS if d in f["name"].lower()), None)
        if not hit:
            skipped.append(f["name"])
            continue
        raw = export(svc, f["id"], XLSX) if f["mimeType"] == GSHEET else download(svc, f["id"])
        try:
            rows, c = parse_gsc(raw, f["name"], hit[1], hit[0])
        except Exception as e:  # file rusak / bukan export GSC
            print("skip", f["name"], e, file=sys.stderr)
            skipped.append(f["name"])
            continue
        gsc += rows
        cut = cut or c
    write("gsc.json", {"rows": gsc, "cut": cut, "skipped": skipped})
    report["gsc"] = len(gsc)

    write("meta.json", {"updated": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"), "rows": report, "gsc_files": len(files) - len(skipped)})
    print(json.dumps(report))


if __name__ == "__main__":
    main()
